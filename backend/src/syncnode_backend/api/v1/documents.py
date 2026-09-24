"""
SyncNode — Document Parser & Knowledge Ingest API.

POST /api/v1/documents/parse    — parse a file (path) and return extracted text
POST /api/v1/documents/ingest   — parse + chunk + embed into ChromaDB RAG
POST /api/v1/documents/upload   — multipart file upload → parse → optional ingest
GET  /api/v1/documents/status   — RAG collection stats

Supported formats (all offline, no cloud):
  .docx  — python-docx
  .xlsx  — openpyxl
  .pptx  — python-pptx
  .pdf   — PyMuPDF (fitz) or pdfplumber fallback
  .txt .md .csv .json .py .ts .js .html — plain text / UTF-8
"""

from __future__ import annotations

import hashlib
import io
import logging
import os
import re
import uuid
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel

logger = logging.getLogger(__name__)
router = APIRouter()


# ── Request / response models ────────────────────────────────────────────────

class ParseRequest(BaseModel):
    path: str
    max_chars: int = 200_000


class ParseResponse(BaseModel):
    path: str
    filename: str
    mime_type: str
    char_count: int
    page_count: Optional[int] = None
    text: str
    sha256: str


class IngestRequest(BaseModel):
    path: str
    title: Optional[str] = None
    tags: list[str] = []
    chunk_size: int = 800
    chunk_overlap: int = 100


class IngestResponse(BaseModel):
    doc_id: str
    path: str
    filename: str
    chunks: int
    collection: str
    sha256: str
    already_existed: bool = False


class RagStatusResponse(BaseModel):
    collection: str
    document_count: int
    chunk_count: int
    persist_dir: str


# ── Parser core ──────────────────────────────────────────────────────────────

_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".yaml", ".yml",
    ".py", ".ts", ".tsx", ".js", ".jsx", ".html", ".htm",
    ".xml", ".toml", ".ini", ".cfg", ".log", ".rst",
}


def _sha256_path(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_text(path: Path, max_chars: int = 200_000) -> tuple[str, str, Optional[int]]:
    """
    Extract plain text from a file.

    Returns:
        (text, mime_type, page_count)
    """
    ext = path.suffix.lower()

    # ── Plain text ──────────────────────────────────────────────────────
    if ext in _TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text[:max_chars], "text/plain", None

    # ── Word .docx ──────────────────────────────────────────────────────
    if ext == ".docx":
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(str(path))
            parts = []
            for para in doc.paragraphs:
                t = para.text.strip()
                if t:
                    parts.append(t)
            # Also extract table cells
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        parts.append(" | ".join(row_texts))
            text = "\n".join(parts)
            return text[:max_chars], "application/vnd.openxmlformats-officedocument.wordprocessingml.document", None
        except Exception as exc:
            raise ValueError(f"Failed to parse .docx: {exc}") from exc

    # ── Excel .xlsx ─────────────────────────────────────────────────────
    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"## Sheet: {sheet_name}")
                for row in ws.iter_rows(values_only=True):
                    row_texts = [str(c) for c in row if c is not None and str(c).strip()]
                    if row_texts:
                        parts.append(" | ".join(row_texts))
            wb.close()
            text = "\n".join(parts)
            return text[:max_chars], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None
        except Exception as exc:
            raise ValueError(f"Failed to parse .xlsx: {exc}") from exc

    # ── PowerPoint .pptx ────────────────────────────────────────────────
    if ext in (".pptx", ".ppt"):
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
            text = "\n".join(parts)
            return text[:max_chars], "application/vnd.openxmlformats-officedocument.presentationml.presentation", len(prs.slides)
        except Exception as exc:
            raise ValueError(f"Failed to parse .pptx: {exc}") from exc

    # ── PDF ─────────────────────────────────────────────────────────────
    if ext == ".pdf":
        # Try PyMuPDF first (faster, better layout)
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(path))
            parts = []
            for page in doc:
                parts.append(page.get_text("text"))
            doc.close()
            text = "\n".join(parts)
            return text[:max_chars], "application/pdf", len(parts)
        except ImportError:
            pass
        # Fallback: pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                parts = [page.extract_text() or "" for page in pdf.pages]
                n = len(pdf.pages)
            text = "\n".join(parts)
            return text[:max_chars], "application/pdf", n
        except ImportError:
            pass
        raise ValueError(
            "PDF parsing requires PyMuPDF or pdfplumber. "
            "Install with: pip install pymupdf  or  pip install pdfplumber"
        )

    raise ValueError(f"Unsupported file type: {ext!r}. Supported: .docx .xlsx .pptx .pdf .txt .md .csv .json and other text files.")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    """Split text into overlapping chunks, preferring paragraph boundaries."""
    if not text.strip():
        return []
    # Split on double newlines first (paragraph boundaries)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if len(current) + len(para) + 2 <= chunk_size:
            current = (current + "\n\n" + para).strip()
        else:
            if current:
                chunks.append(current)
            # If single para > chunk_size, split by sentence
            if len(para) > chunk_size:
                sentences = re.split(r"(?<=[.!?])\s+", para)
                buf = ""
                for sent in sentences:
                    if len(buf) + len(sent) + 1 <= chunk_size:
                        buf = (buf + " " + sent).strip()
                    else:
                        if buf:
                            chunks.append(buf)
                        buf = sent
                if buf:
                    current = buf
                else:
                    current = ""
            else:
                current = para
    if current:
        chunks.append(current)
    # Add overlap: prepend last `overlap` chars of previous chunk
    if overlap > 0 and len(chunks) > 1:
        overlapped = [chunks[0]]
        for i in range(1, len(chunks)):
            tail = chunks[i - 1][-overlap:]
            overlapped.append(tail + " " + chunks[i])
        return overlapped
    return chunks


async def _ingest_to_chroma(
    doc_id: str,
    filename: str,
    text: str,
    sha256: str,
    title: str,
    tags: list[str],
    chunk_size: int,
    chunk_overlap: int,
) -> tuple[int, str, bool]:
    """
    Chunk text and upsert into ChromaDB.

    Returns: (chunk_count, collection_name, already_existed)
    """
    from syncnode_backend.config.settings import settings
    from syncnode_backend.health.chroma_client import get_chroma_client

    client = get_chroma_client()
    collection_name = settings.syncnode_rag_collection

    # Check if this doc already exists (by sha256)
    try:
        col = client.get_collection(collection_name)
        existing = col.get(where={"sha256": sha256}, limit=1)
        if existing and existing.get("ids"):
            return len(existing["ids"]), collection_name, True
    except Exception:
        pass

    # Get or create collection
    col = client.get_or_create_collection(
        name=collection_name,
        metadata={"hnsw:space": "cosine"},
    )

    # Embed chunks with local SentenceTransformer
    from sentence_transformers import SentenceTransformer
    model = SentenceTransformer(settings.syncnode_embedding_model)

    chunks = chunk_text(text, chunk_size, chunk_overlap)
    if not chunks:
        return 0, collection_name, False

    ids = [f"{doc_id}__chunk_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": doc_id,
            "filename": filename,
            "title": title,
            "tags": ",".join(tags),
            "chunk_index": i,
            "sha256": sha256,
        }
        for i in range(len(chunks))
    ]

    embeddings = model.encode(chunks, show_progress_bar=False).tolist()

    col.upsert(
        ids=ids,
        documents=chunks,
        embeddings=embeddings,
        metadatas=metadatas,
    )
    logger.info("[INGEST] %s → %d chunks in %s", filename, len(chunks), collection_name)
    return len(chunks), collection_name, False


# ── Endpoints ────────────────────────────────────────────────────────────────

@router.post("/documents/parse", response_model=ParseResponse)
async def parse_document(req: ParseRequest):
    """
    Parse a file at the given absolute path and return extracted text.
    Does not ingest into RAG.
    """
    path = Path(req.path).expanduser().resolve()
    if not path.exists():
        raise HTTPException(status_code=404, detail=f"File not found: {path}")
    if not path.is_file():
        raise HTTPException(status_code=400, detail=f"Path is not a file: {path}")

    try:
        text, mime_type, page_count = extract_text(path, req.max_chars)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception as exc:
        logger.error("parse_document error: %s", exc)
        raise HTTPException(status_code=500, detail=f"Extraction failed: {exc}")

    sha256 = _sha256_path(path)
    return ParseResponse(
        path=str(path),
        filename=path.name,
        mime_type=mime_type,
        char_count=len(text),
        page_count=page_count,
        text=text,
        sha256=sha256,
    )


@router.post("/documents/ingest", response_model=IngestResponse)
async def ingest_document(req: IngestRequest):
    """
    Parse a file and ingest it into the ChromaDB RAG knowledge base.
    The file must already exist at the given absolute path.
    """
    path = Path(req.path).expanduser().resolve()
    if not path.exists():
        raise HTTPException(status_code=404, detail=f"File not found: {path}")
    if not path.is_file():
        raise HTTPException(status_code=400, detail=f"Not a file: {path}")

    try:
        text, mime_type, _ = extract_text(path)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Extraction failed: {exc}")

    sha256 = _sha256_path(path)
    doc_id = str(uuid.uuid5(uuid.NAMESPACE_URL, sha256))
    title = req.title or path.stem.replace("_", " ").replace("-", " ").title()

    try:
        chunks, collection, already_existed = await _ingest_to_chroma(
            doc_id=doc_id,
            filename=path.name,
            text=text,
            sha256=sha256,
            title=title,
            tags=req.tags,
            chunk_size=req.chunk_size,
            chunk_overlap=req.chunk_overlap,
        )
    except Exception as exc:
        logger.error("ingest_document chroma error: %s", exc)
        raise HTTPException(status_code=500, detail=f"Ingest failed: {exc}")

    return IngestResponse(
        doc_id=doc_id,
        path=str(path),
        filename=path.name,
        chunks=chunks,
        collection=collection,
        sha256=sha256,
        already_existed=already_existed,
    )


@router.post("/documents/upload", response_model=IngestResponse)
async def upload_and_ingest(
    file: UploadFile = File(...),
    ingest: bool = True,
    chunk_size: int = 800,
    chunk_overlap: int = 100,
):
    """
    Upload a file, save it to the workspace, parse it, and optionally ingest into RAG.
    Supports: .docx .xlsx .pptx .pdf .txt .md .csv .json and other text files.
    """
    from syncnode_backend.config.settings import settings

    data = await file.read()
    if len(data) == 0:
        raise HTTPException(status_code=400, detail="Empty file")

    filename = file.filename or "upload.bin"
    upload_dir = settings.syncnode_workspace_root / "uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    dest = upload_dir / filename

    # Write to disk
    with open(dest, "wb") as f:
        f.write(data)

    sha256 = _sha256_bytes(data)

    if not ingest:
        return IngestResponse(
            doc_id=str(uuid.uuid5(uuid.NAMESPACE_URL, sha256)),
            path=str(dest),
            filename=filename,
            chunks=0,
            collection="",
            sha256=sha256,
        )

    # Parse
    try:
        text, mime_type, _ = extract_text(dest)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Extraction failed: {exc}")

    doc_id = str(uuid.uuid5(uuid.NAMESPACE_URL, sha256))
    title = Path(filename).stem.replace("_", " ").replace("-", " ").title()

    try:
        chunks, collection, already_existed = await _ingest_to_chroma(
            doc_id=doc_id,
            filename=filename,
            text=text,
            sha256=sha256,
            title=title,
            tags=[],
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )
    except Exception as exc:
        logger.error("upload_and_ingest chroma error: %s", exc)
        raise HTTPException(status_code=500, detail=f"Ingest failed: {exc}")

    return IngestResponse(
        doc_id=doc_id,
        path=str(dest),
        filename=filename,
        chunks=chunks,
        collection=collection,
        sha256=sha256,
        already_existed=already_existed,
    )


@router.get("/documents/status", response_model=RagStatusResponse)
async def rag_status():
    """Return ChromaDB collection stats."""
    from syncnode_backend.config.settings import settings
    from syncnode_backend.health.chroma_client import get_chroma_client
    try:
        client = get_chroma_client()
        collection_name = settings.syncnode_rag_collection
        try:
            col = client.get_collection(collection_name)
            all_items = col.get()
            chunk_count = len(all_items.get("ids", []))
            doc_ids = set()
            for meta in (all_items.get("metadatas") or []):
                if meta and meta.get("doc_id"):
                    doc_ids.add(meta["doc_id"])
            doc_count = len(doc_ids)
        except Exception:
            chunk_count = 0
            doc_count = 0
        return RagStatusResponse(
            collection=collection_name,
            document_count=doc_count,
            chunk_count=chunk_count,
            persist_dir=settings.chroma_persist_dir,
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))
